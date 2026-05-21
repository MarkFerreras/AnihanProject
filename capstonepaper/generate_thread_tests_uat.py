"""Generate non-technical (UAT-friendly) Thread Testing Cases for Anihan SRMS."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Thread Testing Cases (UAT)"

headers = ["Test ID", "Title", "Pre-Conditions", "Test Steps", "Expected Results",
           "Tester", "Test Date", "Results", "Comments"]
ws.append(headers)

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
fill_zone_fill = PatternFill(start_color="FFF4D6", end_color="FFF4D6", fill_type="solid")
thin = Side(border_style="thin", color="999999")
border = Border(top=thin, left=thin, right=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")

for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align
    c.border = border

# Each case: (Title, Pre-Conditions, Test Steps, Expected Results) — all plain language
cases = [
    # ---------- AUTH / LOGIN ----------
    ("Admin can log in successfully",
     "The system is running.\nYou have the admin account ready (username: admin, password: password123).\nYou are not currently logged in.",
     "1. Open the login page in your browser.\n2. Type 'admin' in the username box.\n3. Type 'password123' in the password box.\n4. Click the Login button.",
     "You should be brought to the Admin home page.\nYour name should appear in the top-right corner.\nThe system should remember that you logged in (it will show up in the activity logs)."),

    ("Registrar can log in successfully",
     "The system is running.\nThe registrar account is ready (username: registrar, password: password123).\nYou are not currently logged in.",
     "1. Open the login page.\n2. Type 'registrar' in the username box.\n3. Type 'password123' in the password box.\n4. Click Login.",
     "You should land on the Registrar home page.\nThe top menu should show: Home, Subjects, Classes, Sections.\nThe login should be recorded in the activity logs."),

    ("Trainer can log in successfully",
     "The trainer account exists (username: trainer, password: password123).\nYou are logged out.",
     "1. Open the login page.\n2. Type 'trainer' in the username box.\n3. Type 'password123' in the password box.\n4. Click Login.",
     "You should land on the Trainer home page.\nThe top menu should show: Home, My Subjects, My Classes."),

    ("Wrong password is rejected at login",
     "You know 'admin' is a valid username.\nYou will deliberately type a wrong password.",
     "1. Open the login page.\n2. Type 'admin' in the username box.\n3. Type 'wrongpass' in the password box.\n4. Click Login.",
     "Login should fail.\nA message should appear saying the username or password is wrong.\nYou should still be on the login page."),

    ("Disabled (deactivated) account cannot log in",
     "There is a user account that an admin previously deactivated.",
     "1. Open the login page.\n2. Type the deactivated user's username and password.\n3. Click Login.",
     "Login should fail.\nThe system should not allow this person to enter, even if their password is correct."),

    ("Log out works correctly",
     "You are logged in to any account.",
     "1. Click your name/icon at the top-right of the page.\n2. Click 'Log Out'.\n3. Confirm if asked.",
     "You should be sent back to the login page.\nYour session should end (going 'back' in the browser should not let you back in).\nThe logout should be recorded in the activity logs."),

    ("Inactive sessions automatically expire",
     "You are logged in.\nThe system is set to end inactive sessions after 30 minutes.",
     "1. Log in as any user.\n2. Leave the browser open and do not click anything for 30+ minutes.\n3. After waiting, click any menu link or refresh the page.",
     "The system should automatically take you back to the login page.\nYou should not be able to see any private information until you log in again."),

    ("People can only see pages meant for their role",
     "You are logged in as a trainer.\nThe admin pages are only meant for admins.",
     "1. While logged in as trainer, try to open the admin home page directly by typing its address.",
     "You should not be allowed in.\nThe system should send you back to your own page or to the login page.\nAdmin information should never appear."),

    # ---------- ADMIN DASHBOARD ----------
    ("Admin sees the list of all users",
     "You are logged in as admin.\nThere are user accounts already in the system.",
     "1. Open the Admin home page.\n2. Wait for the users table to load.",
     "A table of users should appear showing each person's Last Name, First Name, Username, Role, and Status.\nEach row should have buttons for Edit and Delete."),

    ("Admin can create a brand-new user",
     "You are logged in as admin.\nThe username you want to create has never been used before.",
     "1. On the Admin page, click 'Add User'.\n2. Fill in last name, first name, a new username, a strong password, role (e.g., Trainer), and birthdate.\n3. Click Save.",
     "The new user should appear in the table.\nTheir age should be filled in automatically from the birthdate (you don't have to type it).\nThe action should be recorded in the activity logs."),

    ("System blocks duplicate usernames",
     "You are logged in as admin.\nThere is already a user with the username 'admin' in the system.",
     "1. Click 'Add User'.\n2. Type 'admin' as the username and fill in all other fields.\n3. Click Save.",
     "The system should refuse to save.\nA message should appear saying the username is already taken.\nThe table should stay the same."),

    ("Admin can edit a user's personal details",
     "You are logged in as admin.\nA user (not yourself) exists in the table.",
     "1. Click 'Edit' next to a user.\n2. Change their last name and birthdate.\n3. Click Save.",
     "The user's details should update in the table.\nTheir age should automatically refresh from the new birthdate.\nThe change should be recorded in the activity logs."),

    ("Admin can reset another user's password",
     "You are logged in as admin.\nThe user whose password you want to reset exists.",
     "1. Click 'Edit' on the user.\n2. Type a new password (at least 8 characters).\n3. Click Save.",
     "The password should be saved.\nThe user should be able to log in with the new password.\nTheir old password should no longer work."),

    ("Admin cannot accidentally change their own role",
     "You are logged in as admin.\nYou are editing your own account.",
     "1. Open the Edit screen for your own admin account.\n2. Try to change your role from Admin to Registrar.\n3. Click Save.",
     "The system should refuse the change.\nYou should remain an admin after the attempt."),

    ("Admin can deactivate a user (soft delete)",
     "You are logged in as admin.\nThe user you want to deactivate is currently active.",
     "1. Click 'Delete' on a user.\n2. Choose 'Deactivate Account' from the choices.\n3. Confirm.",
     "The user's status should change to 'Disabled' in the table.\nThe deactivated user should no longer be able to log in.\nThey can be re-enabled later if needed."),

    ("Admin can permanently delete a user (with safety check)",
     "You are logged in as admin.\nThe user to delete is no longer needed at all.",
     "1. Click 'Delete' on a user.\n2. Choose 'Permanently Delete'.\n3. In the warning box, type the word 'delete' in the text field.\n4. Click the confirm button.",
     "The confirm button should only become active after you type 'delete' correctly.\nThe user should disappear from the table.\nThe action should be recorded in the activity logs."),

    ("Admin can re-enable a deactivated user",
     "You are logged in as admin.\nA user is currently deactivated.",
     "1. Find the deactivated user in the table.\n2. Click the re-enable button.\n3. Confirm.",
     "The user's status should change back to Active.\nThey should be able to log in again."),

    # ---------- SYSTEM LOGS ----------
    ("Activity logs page shows the last 7 days by default",
     "You are logged in as admin.\nThere are activity log entries in the system.",
     "1. Open the Logs page from the menu.",
     "A table of activities should appear, showing only entries from the last 7 days.\nEach row should show who did what, when, and from where."),

    ("Admin can filter logs by a custom date range",
     "You are logged in as admin.\nThere are log entries spread across many dates.",
     "1. On the Logs page, choose the 'Custom Range' option.\n2. Pick a start date and an end date covering a 3-day window.\n3. Click Apply.",
     "Only the entries within those dates should be shown.\nEntries outside the range should not appear."),

    ("Admin can export logs as a CSV file",
     "You are on the Logs page.",
     "1. Click the 'Export' button.\n2. Choose 'CSV'.",
     "A CSV file should download.\nWhen opened, it should contain the same log entries that were on screen."),

    ("Admin can export logs as an Excel file",
     "You are on the Logs page.",
     "1. Click 'Export'.\n2. Choose 'Excel (XLSX)'.",
     "An Excel file should download.\nIt should contain the filtered log entries with proper column headers."),

    ("Admin can export logs as a Word document",
     "You are on the Logs page.",
     "1. Click 'Export'.\n2. Choose 'Word (DOCX)'.",
     "A Word document should download.\nIt should contain a neatly formatted table of the log entries."),

    # ---------- EDIT ACCOUNT MODAL ----------
    ("User can update their own personal details",
     "You are logged in (as any role).",
     "1. Click your name/icon at the top.\n2. Click 'Edit Account'.\n3. Open the 'Personal Details' tab.\n4. Change your last name, first name, or birthdate.\n5. Click Save.",
     "Your details should be saved.\nYour age should update automatically as you change the birthdate.\nWhen you reopen the page, your new details should still be there."),

    ("User can change their own username",
     "You are logged in.\nThe new username you want is not already taken.",
     "1. Open 'Edit Account' from your top menu.\n2. Go to 'Account Settings'.\n3. Type a new username.\n4. Type your current password to confirm.\n5. Click Save.",
     "Your username should be updated.\nThe new username should appear on the page.\nNext time you log in, you must use the new username."),

    ("User can change their own password",
     "You are logged in.\nYou know your current password.",
     "1. Open 'Edit Account' > 'Account Settings'.\n2. Type your current password.\n3. Type a new strong password (at least 8 characters, with uppercase, lowercase, number, and special character).\n4. Click Save.",
     "Your password should be updated.\nNext time you log in, you must use the new password.\nThe old password should no longer work."),

    ("System rejects weak passwords when you change your own",
     "You are logged in.\nYou are about to type a weak password on purpose.",
     "1. Open 'Edit Account' > 'Account Settings'.\n2. Type a weak password like 'short' as the new password.\n3. Click Save.",
     "The system should refuse to save the password.\nA message should explain that the password does not meet the security rules.\nYour current password should still be in effect."),

    ("Show/hide password eye-icon works",
     "You are on any page that has a password field.",
     "1. Click in the password field.\n2. Type a few letters.\n3. Click the eye icon at the right side of the field.\n4. Click the eye icon again.",
     "When you click the eye, the password should become visible.\nWhen you click again, it should hide back as dots."),

    # ---------- REGISTRAR DASHBOARD ----------
    ("Registrar sees the list of student records",
     "You are logged in as registrar.\nThere are student records in the system.",
     "1. Open the Registrar home page.",
     "A table of students should appear with columns like Student ID, Name, Status, Batch, Course, and Section.\nA search box and a status filter should be visible above the table."),

    ("Registrar can search for students by name",
     "You are on the Registrar home page.\nAt least one student exists with a known last name.",
     "1. Type part of a student's last name into the search box.",
     "The table should refresh and show only students whose names (or other key information) match what you typed."),

    ("Registrar can filter students by batch year range",
     "You are on the Registrar home page.\nStudents exist across different batch years.",
     "1. Type a 'From Year' (e.g., 2024).\n2. Type a 'To Year' (e.g., 2026).\n3. Click Apply.",
     "Only students whose batch year falls inside the chosen range should be shown."),

    ("Registrar can filter students by status",
     "You are on the Registrar home page.\nStudents exist with different statuses (Active, Submitted, Enrolling, Graduated).",
     "1. Click the status dropdown.\n2. Choose 'Active'.",
     "Only students whose status is 'Active' should appear in the table."),

    ("Registrar can view a student's full details",
     "You are on the Registrar home page.\nAt least one student record exists.",
     "1. Click 'Open Details' on a student row.",
     "A window should pop up showing the student's complete personal details, parents, and guardian.\nA colored badge should show their status (green for Active, grey for Submitted, blue for Graduated)."),

    # ---------- EDIT STUDENT RECORD ----------
    ("Registrar can update a student's full record",
     "You are logged in as registrar.\nA student record exists.",
     "1. From the registrar page, click Edit on a record.\n2. Change details like last name, birthdate, parents' information.\n3. Click Save.",
     "All changes should be saved.\nThe student's age should automatically update from the new birthdate.\nThe change should be recorded in the activity logs."),

    ("Registrar can update a student's OJT information",
     "You are editing a student record.",
     "1. Fill in the OJT company name, address, and hours.\n2. Click Save.",
     "The OJT details should be saved.\nIf the student already had OJT info, it should be updated; if not, a new OJT record should be created."),

    ("Registrar can add and remove TESDA qualifications",
     "You are editing a student record.",
     "1. Fill in TESDA qualification details in slots 1 and 2.\n2. Click Save.\n3. Open the record again, clear slot 2 and fill slot 3.\n4. Click Save.",
     "The TESDA list should match exactly what was on screen when you saved.\nThe system should not complain about duplicates."),

    ("Registrar can add and delete School Year rows",
     "You are editing a student record.",
     "1. Click 'Add Row' on the School Years table three times.\n2. Fill in each row.\n3. Click Save.\n4. Open the record again, delete one row, and save.",
     "The saved school years should match exactly what was on screen."),

    ("System warns before leaving with unsaved changes",
     "You are editing a student record with changes that have not been saved yet.",
     "1. Type or change something on the form.\n2. Try to navigate to another page by clicking a menu link or pressing the back button.",
     "A warning message should appear asking if you really want to leave with unsaved changes.\nYou should be able to either stay or leave."),

    ("New student is automatically placed in the current year's batch",
     "You are logged in as registrar.\nA batch for the current year exists in the system.",
     "1. Save a student record that does not have a batch assigned yet.",
     "The student should be automatically placed in the current year's batch.\nNo extra step from the registrar should be needed."),

    ("Registrar can delete a student record with safety check",
     "You are logged in as registrar.\nA student record exists (with parents, guardian, uploads, etc.).",
     "1. Open the student's details.\n2. Click 'Delete'.\n3. In the warning box, type the word 'delete'.\n4. Click confirm.",
     "The confirm button should only activate when you type 'delete' correctly.\nThe student record AND all their related information (uploads, parents, guardian, education, grades, etc.) should be removed.\nUploaded files should also be removed from the server."),

    # ---------- SUBJECTS ----------
    ("Registrar can create a new subject",
     "You are logged in as registrar.\nAt least one qualification exists in the system.",
     "1. Open the Subjects page.\n2. Click 'Create Subject'.\n3. Type the subject code, subject name, choose a qualification, and set the units.\n4. Click Save.",
     "The new subject should appear in the table.\nThe action should be recorded in the activity logs."),

    ("Registrar can edit an existing subject",
     "You are on the Subjects page.\nA subject exists.",
     "1. Click 'Edit' on a subject.\n2. Change the name and units (subject code cannot be changed).\n3. Click Save.",
     "The updated subject details should appear in the table.\nThe subject code should stay the same."),

    ("Registrar can assign a default trainer to a subject",
     "You are on the Subjects page.\nThere are trainer accounts available.",
     "1. Click 'Assign Trainer' on a subject.\n2. Pick a trainer from the dropdown.\n3. Click Save.",
     "The trainer's name should appear on the subject's row.\nThe action should be recorded in the activity logs."),

    ("System blocks deleting a subject that is still in use",
     "You are on the Subjects page.\nThe subject you want to delete is used in a class or already has grades.",
     "1. Click Delete on that subject.\n2. Type 'delete' to confirm.\n3. Click confirm.",
     "The system should refuse to delete.\nA friendly message should explain that the subject is still being used in classes or grades.\nThe subject should remain in the table."),

    ("Registrar can delete a subject that is not in use",
     "You are on the Subjects page.\nThe subject you want to delete has no classes and no grades attached.",
     "1. Click Delete on the subject.\n2. Type 'delete' to confirm.\n3. Click confirm.",
     "The subject should disappear from the table.\nThe action should be recorded in the activity logs."),

    # ---------- CLASSES ----------
    ("Registrar can create a new class",
     "You are logged in as registrar.\nA section and a subject exist (and optionally a trainer).",
     "1. Open the Classes page.\n2. Click 'Create Class'.\n3. Choose a section and a subject (a default trainer may auto-fill from the subject).\n4. Confirm the semester.\n5. Click Save.",
     "The new class should appear in the table with 0 enrolled students.\nThe system should prevent two identical classes (same section, subject, and semester) from existing at the same time."),

    ("System prevents creating duplicate classes",
     "There is already a class for the same section, subject, and semester.",
     "1. Try to create another class with exactly the same section, subject, and semester.",
     "The system should refuse to save.\nA message should explain that this class already exists."),

    ("Registrar can change a class's assigned trainer",
     "You are on the Classes page.\nA class exists with a trainer assigned.",
     "1. Click 'Edit Trainer' on the class.\n2. Pick a different trainer (or choose 'Unassigned').\n3. Click Save.",
     "The new trainer's name should appear on the class row.\nIf you chose 'Unassigned', the row should show 'Unassigned' in italics.\nThe change should be recorded in the activity logs."),

    ("Registrar can enroll a single student into a class",
     "A class exists.\nThere is an eligible student (Active or Submitted, in the same section, not yet enrolled).",
     "1. Open the 'Manage Students' window for the class.\n2. Pick an eligible student from the list.\n3. Click 'Enroll'.",
     "The student should appear in the class's roster.\nThe enrolled count for the class should increase by 1."),

    ("Registrar can enroll an entire section into a class at once",
     "A class exists.\nA section has several Active students; some may already be enrolled.",
     "1. Open 'Manage Students' on the class.\n2. Click 'Enroll Whole Section'.",
     "A summary should appear telling you how many were enrolled, how many were skipped because they are already enrolled, and how many were skipped because they were not eligible.\nThe action should be recorded in the activity logs."),

    ("Registrar can remove a student from a class",
     "A class has at least one enrolled student.",
     "1. Open 'Manage Students' on the class.\n2. Click 'Remove' next to a student.\n3. Confirm.",
     "The student should be removed from the class.\nThe enrolled count should decrease by 1.\nThe action should be recorded in the activity logs."),

    ("Class page shows the current semester automatically",
     "You are logged in as registrar.\nAt least one batch exists.",
     "1. Open the Classes page.",
     "The current semester (based on the latest batch year) should be shown.\nIt should also be pre-filled when creating a new class."),

    # ---------- SECTIONS ----------
    ("Registrar can create a new section",
     "You are logged in as registrar.\nA batch and a course exist.",
     "1. Open the Sections page.\n2. Click 'Create Section'.\n3. Type the section code and section name; choose a batch and a course.\n4. Click Save.",
     "The new section should appear in the table.\nThe action should be recorded in the activity logs."),

    ("Registrar can rename a section",
     "A section exists.",
     "1. Click 'Edit' on a section.\n2. Change the section name.\n3. Click Save.",
     "The new section name should appear in the table."),

    ("Registrar can assign students to a section",
     "A section exists.\nThere are eligible students (not yet in any section, with status Submitted).",
     "1. Click 'Manage Students' on a section.\n2. Open the 'Add Students' tab.\n3. Filter by batch and/or course.\n4. Tick the students you want and click 'Assign'.",
     "A summary should show how many were successfully assigned and which (if any) were skipped.\nThe students' status should change from 'Submitted' to 'Active'."),

    ("Registrar can remove a student from a section",
     "A section has one or more students in it.",
     "1. Click 'Manage Students' > 'Current Students' tab.\n2. Click 'Remove' on a student.\n3. Confirm.",
     "The student should no longer appear in the section's roster.\nTheir status should go back from 'Active' to 'Submitted'.\nAny class enrollments tied to the section should be removed automatically."),

    ("System blocks deleting a section that still has classes",
     "A section is being used by one or more classes.",
     "1. Click 'Delete' on the section.\n2. Type 'delete' to confirm.\n3. Click confirm.",
     "The system should refuse to delete.\nA clear message should say: classes still reference this section; remove those classes first.\nThe section should remain."),

    ("Registrar can delete an unused section",
     "A section has no classes and no students assigned.",
     "1. Click 'Delete' on the section.\n2. Type 'delete' to confirm.\n3. Click confirm.",
     "The section should disappear from the table.\nThe action should be recorded in the activity logs."),

    # ---------- TRAINER PAGES ----------
    ("Trainer home page loads correctly",
     "You are logged in as trainer.",
     "1. Open the Trainer home page.",
     "A welcome section should appear with quick-link cards.\nThe top menu should show: Home, My Subjects, My Classes."),

    ("Trainer can see their assigned subjects and student rosters",
     "The trainer is assigned to at least one class.",
     "1. Open 'My Subjects'.\n2. Click on a subject.",
     "A table of the trainer's subjects should appear, showing the number of enrolled students and the sections/courses involved.\nClicking a subject should expand a list of students under it."),

    ("Trainer can see their assigned classes and student rosters",
     "The trainer is assigned to at least one class.",
     "1. Open 'My Classes'.\n2. Click on a class.",
     "A table of the trainer's classes should appear.\nClicking a class should expand a list of the students enrolled in that class."),

    ("Trainer cannot view classes that don't belong to them",
     "You are logged in as a trainer.\nA class exists that is owned by a DIFFERENT trainer.",
     "1. Try to view the other trainer's class roster (for example, by typing the class address directly).",
     "The system should refuse to show the information.\nNo student names should appear."),

    # ---------- GRADE INPUT ----------
    ("Trainer can open the Grade Input window",
     "You are logged in as a trainer who owns at least one class with enrolled students.",
     "1. On 'My Classes', click 'Input Grades' on a class.",
     "A window should appear listing all enrolled students with empty grade fields ready to be filled in."),

    ("Trainer can save grades for students",
     "The Grade Input window is open.\nAt least one student row will be filled in.",
     "1. Type a midterm grade and a finals grade for one or more students.\n2. Click 'Save'.",
     "The grades should be saved.\nThe final grade should be calculated and shown automatically.\nA confirmation message should appear."),

    ("System warns when trying to save with no grades filled",
     "The Grade Input window is open.\nNo grades have been entered yet.",
     "1. Click 'Save' without typing any grade.",
     "A warning should pop up saying no grades were entered.\nNothing should be saved."),

    ("Trainer can lock grades to prevent further edits",
     "The trainer has already saved grades for a class.",
     "1. Click 'Lock Grades'.\n2. Confirm.",
     "The grade fields should become read-only.\nA 'Locked' indicator should be visible.\nThe action should be recorded in the activity logs."),

    ("Trainer can unlock grades to allow editing again",
     "The grades for a class are currently locked.",
     "1. Click 'Unlock Grades'.\n2. Confirm.",
     "The grade fields should become editable again.\nThe lock indicator should go away."),

    # ---------- STUDENT PORTAL ----------
    ("Student Portal welcome page is public (no login needed)",
     "You are not logged in.",
     "1. Open the Student Portal page in a browser.",
     "The welcome page should appear with name fields.\nThere should be no login prompt and no account menu."),

    ("System warns when a duplicate student name is entered",
     "A student with a known name (e.g., 'Dela Cruz, Maria, Santos') is already in the system.",
     "1. On the Student Portal, type the same Last Name, First Name, and Middle Name.\n2. Click Submit.",
     "A warning should appear saying a student with that name already exists.\nThe portal should not let you proceed to the next page."),

    ("New student can proceed to the details form",
     "No existing student has the same name as the one being entered.",
     "1. Type a unique Last Name, First Name, and Middle Name.\n2. Click Submit.",
     "You should be brought to the Student Details form (the enrollment wizard).\nThe name fields should already be filled in."),

    ("Form blocks moving forward if required fields are empty",
     "You are on Step 1 of the Student Details form.",
     "1. Leave required fields blank (for example, Civil Status).\n2. Click Next.",
     "A red highlight or warning should appear on the empty fields.\nThe form should not move to the next step until they are filled in."),

    ("Baptismal certificate is optional",
     "You are on Step 2 of the Student Details form.",
     "1. Check the 'Baptized' box.\n2. Fill in the Baptism Date and Place.\n3. Do NOT upload a baptismal certificate file.\n4. Click Next.",
     "The form should let you continue to the next step.\nNo error should appear about a missing certificate."),

    ("Education table has 4 columns",
     "You are on Step 3 (Educational Background) of the Student Details form.",
     "1. Look at the education table.\n2. Fill in 4 rows (School Name, Address, Course, School Year).",
     "The table should have only 4 columns (no Grade/Year, no Semester column).\nYou should be able to type into each row and continue."),

    ("Uploaded files are only saved after the form is submitted",
     "You are on the Student Details form.\nID Photo and other file fields are optional.",
     "1. Choose a file for the ID Photo (you should see a small preview).\n2. Continue through the rest of the wizard.\n3. On the last step, click Submit.",
     "The text details should be saved first.\nOnly after the student record is saved should the file actually be uploaded.\nA message like 'Uploaded: <filename>' should appear."),

    ("Final submit saves everything together",
     "All steps of the Student Details form are completed and valid.",
     "1. Click 'Submit' on the final step.",
     "All the information (student details, parents, guardian, education, school years) should be saved together.\nThe new student's status should be set to 'Submitted'.\nA confirmation message should appear."),

    ("System blocks resubmitting the same student details",
     "The student has already submitted their details.",
     "1. Try to submit the same student details again.",
     "The system should refuse.\nA message should explain that this student has already submitted their information.\nNo duplicate records should be created."),

    # ---------- NON-FUNCTIONAL / CROSS-CUTTING ----------
    ("Unexpected errors show a friendly message (not technical jargon)",
     "You are logged in.\nAn unexpected server error is forced (for testing).",
     "1. Trigger an unexpected error situation.",
     "The user should see a friendly message like 'An unexpected error occurred. Please contact the administrator.'\nNo technical details (like database table names, SQL, or programmer text) should appear on screen."),

    ("Data conflict errors show a clear message",
     "A data conflict is forced (e.g., trying to delete something that is still being used).",
     "1. Trigger a data conflict.",
     "A clear, friendly message should appear explaining the conflict.\nNo technical or database-level details should be shown to the user."),

    ("Activity logs cannot be edited or deleted",
     "Activity log entries exist in the system.",
     "1. Try (from inside the application) to edit or delete a log entry.",
     "There should be no option anywhere in the system to change or remove activity log entries.\nThe logs should be a permanent record of what happened."),

    ("Login session is protected (security check)",
     "You are logged in.",
     "1. Have a developer/IT person check your browser's session cookie settings.",
     "The session cookie should be marked as HttpOnly (cannot be read by webpage scripts).\nIt should expire after 30 minutes of inactivity.\nThis keeps your login safe from common attacks."),

    ("The system can handle many users at the same time",
     "The system is set up like the real working environment.\nTarget: 50 or more users at once.",
     "1. Have 50 different users log in at the same time and do their normal tasks (add students, view records, input grades).",
     "Everyone should be able to use the system at the same time without it slowing down badly or crashing.\nNo one's information should accidentally appear in someone else's session."),

    ("Document uploads finish in under 3 seconds",
     "You are logged in as registrar.\nYou have a document file smaller than 5 MB ready to upload.",
     "1. Upload a document for a student.",
     "The upload should finish within 3 seconds.\nThe file should appear attached to the correct student afterwards."),

    ("Student records page loads quickly even with many records",
     "There are at least 200 student records in the system.",
     "1. Open the Registrar home page.",
     "The student table should appear within 5 seconds.\nSearch and filter should also respond quickly."),
]

# Write rows
for i, (title, pre, steps, exp) in enumerate(cases, start=1):
    test_id = f"TC-{i:03d}"
    ws.append([test_id, title, pre, steps, exp, "", "", "", ""])

# Column widths
ws.column_dimensions['A'].width = 10  # Test ID
ws.column_dimensions['B'].width = 45  # Title
ws.column_dimensions['C'].width = 48  # Pre-Conditions
ws.column_dimensions['D'].width = 55  # Test Steps
ws.column_dimensions['E'].width = 60  # Expected Results
ws.column_dimensions['F'].width = 18  # Tester
ws.column_dimensions['G'].width = 14  # Test Date
ws.column_dimensions['H'].width = 14  # Results
ws.column_dimensions['I'].width = 35  # Comments

# Apply wrap, border, alternating fill (cols A-E), fill-zone tint (cols F-I)
alt_fill = PatternFill(start_color="F4F7FB", end_color="F4F7FB", fill_type="solid")
fill_zone_alt = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
for row in range(2, ws.max_row + 1):
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=row, column=col)
        c.alignment = wrap
        c.border = border
        if col >= 6:  # tester/date/results/comments — fill-in zone
            c.fill = fill_zone_alt if row % 2 == 0 else fill_zone_fill
        else:
            if row % 2 == 0:
                c.fill = alt_fill

# Bold Test ID column
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=1).font = Font(bold=True)

# Data validation (dropdown) on Results column (H)
dv = DataValidation(type="list",
                    formula1='"Pass,Fail,Blocked,Not Tested"',
                    allow_blank=True,
                    showDropDown=False)
dv.error = "Please choose Pass, Fail, Blocked, or Not Tested."
dv.errorTitle = "Invalid value"
dv.prompt = "Choose one: Pass / Fail / Blocked / Not Tested"
dv.promptTitle = "Test Result"
ws.add_data_validation(dv)
dv.add(f"H2:H{ws.max_row}")

# Date format hint on column G
from openpyxl.styles import NamedStyle
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=7).number_format = "yyyy-mm-dd"

# Freeze header
ws.freeze_panes = "A2"

out_path = r"c:\Users\emman\Desktop\Capstone Files\DEVELOPMENT\AnihanProject\capstonepaper\Anihan_SRMS_Thread_Testing_Cases_UAT.xlsx"
wb.save(out_path)
print(f"Saved: {out_path}")
print(f"Total test cases: {len(cases)}")
